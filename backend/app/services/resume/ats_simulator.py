from io import BytesIO
from pydantic import BaseModel
from typing import List
import pdfplumber
from docx import Document as DocxDocument
from app.schemas.resume_v2 import FormatIssue


class ATSResult(BaseModel):
    parseability_score: int
    raw_text: str
    format_issues: List[FormatIssue]


def simulate_ats(content: bytes, filename: str) -> ATSResult:
    ext = filename.lower().rsplit(".", 1)[-1]
    if ext == "pdf":
        return _simulate_pdf(content)
    if ext == "docx":
        return _simulate_docx(content)
    raise ValueError(f"Unsupported: {ext}")


def _simulate_pdf(content: bytes) -> ATSResult:
    issues: List[FormatIssue] = []
    score = 100
    with pdfplumber.open(BytesIO(content)) as pdf:
        text_parts: list[str] = []
        for i, page in enumerate(pdf.pages):
            tables = page.find_tables()
            if tables:
                issues.append(FormatIssue(
                    type="table",
                    location=f"page {i + 1}",
                    fix_hint="ATS parsers misread tables. Use bullet lists instead.",
                ))
                score -= 15
            if page.images:
                issues.append(FormatIssue(
                    type="image",
                    location=f"page {i + 1}",
                    fix_hint="Remove images. ATS ignores them but they shift layout.",
                ))
                score -= 5
            text = page.extract_text() or ""
            text_parts.append(text)
            cols = _detect_columns(page)
            if cols > 1:
                issues.append(FormatIssue(
                    type="multi_column",
                    location=f"page {i + 1}",
                    fix_hint="Switch to single-column layout.",
                ))
                score -= 20
    raw = "\n".join(text_parts)
    score = max(0, min(100, score))
    return ATSResult(parseability_score=score, raw_text=raw, format_issues=issues)


def _detect_columns(page) -> int:
    """Detect repeated side-by-side text bands separated by a stable gutter."""
    return 2 if _column_confidence(page) >= 0.65 else 1


def _column_confidence(page) -> float:
    words = page.extract_words()
    if not words:
        return 0.0
    width = page.width
    min_gutter = width * 0.18
    min_side_words = 2
    line_tolerance = 3

    lines: dict[int, list[dict]] = {}
    for word in words:
        key = round(float(word["top"]) / line_tolerance)
        lines.setdefault(key, []).append(word)

    split_lines = 0
    candidate_lines = 0
    gutters: list[tuple[float, float]] = []
    for line_words in lines.values():
        if len(line_words) < min_side_words * 2:
            continue
        candidate_lines += 1
        ordered = sorted(line_words, key=lambda w: float(w["x0"]))
        gaps = [
            (float(right["x0"]) - float(left["x1"]), float(left["x1"]), float(right["x0"]))
            for left, right in zip(ordered, ordered[1:])
        ]
        if not gaps:
            continue
        gap, left_edge, right_edge = max(gaps, key=lambda item: item[0])
        if gap < min_gutter:
            continue
        left_count = sum(1 for word in ordered if float(word["x1"]) <= left_edge)
        right_count = sum(1 for word in ordered if float(word["x0"]) >= right_edge)
        if left_count >= min_side_words and right_count >= min_side_words:
            split_lines += 1
            gutters.append((left_edge, right_edge))

    if candidate_lines < 6 or split_lines < 4:
        return 0.0
    if not gutters:
        return 0.0

    gutter_centers = [(left + right) / 2 for left, right in gutters]
    center_avg = sum(gutter_centers) / len(gutter_centers)
    spread = max(abs(center - center_avg) for center in gutter_centers)
    stability = max(0.0, 1.0 - spread / (width * 0.15))
    prevalence = split_lines / candidate_lines
    return min(1.0, prevalence * 0.75 + stability * 0.25)


def _simulate_docx(content: bytes) -> ATSResult:
    doc = DocxDocument(BytesIO(content))
    issues: List[FormatIssue] = []
    score = 100
    if doc.tables:
        issues.append(FormatIssue(
            type="table",
            location="document",
            fix_hint="ATS parsers misread tables. Use bullet lists.",
        ))
        score -= 15
    raw = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    return ATSResult(parseability_score=max(0, min(100, score)), raw_text=raw, format_issues=issues)
